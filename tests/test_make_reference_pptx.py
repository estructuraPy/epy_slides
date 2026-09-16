"""Tests for the ``make_reference_pptx`` dev-tooling script.

Covers the pure hex/font helpers, each private XML-mutating helper
against a real python-pptx object (built from ``Presentation()``'s
bundled default template), the ``build_reference`` orchestration end
to end (including its ``ValueError`` guard), and ``main`` (including
its ``if __name__ == "__main__":`` guard) — all without ever writing
into the real ``src/`` asset tree.
"""

from __future__ import annotations

import runpy
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from epy_slides._core._packaging import make_reference_pptx as mrp
from epy_slides._core.themes_base import Theme

# --------------------------------------------------------------- helpers


def _make_theme(theme_id: str = "test-theme") -> Theme:
    """Build a ``Theme`` with every ``_COLOR_MAP`` target populated."""
    return Theme(
        id=theme_id,
        display_name="Test Theme",
        qt_palette={},
        css_vars={
            "fg": "#111111",
            "bg": "#eeeeee",
            "heading-color": "#222222",
            "bg-soft": "#dddddd",
            "link": "#3366ff",
            "link-hover": "#ff6633",
            "mark-bg": "#ffff00",
            "table-header-bg": "#cccccc",
            "border": "#999999",
            "quote-rule": "#00aa00",
            "font-family-headings": '"Header Font", sans-serif',
            "font-family-text": '"Body Font", serif',
        },
    )


def _fresh_theme_elements():
    """Return a fresh, disconnected ``(clrScheme, fontScheme)`` pair.

    Parsed the same way ``build_reference`` does: from a brand-new
    ``Presentation()``'s theme part blob, so mutating them can never
    leak between tests.
    """
    prs = Presentation()
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    theme_el = etree.fromstring(theme_part.blob)
    elements = theme_el.find(qn("a:themeElements"))
    clr_scheme = elements.find(qn("a:clrScheme"))
    font_scheme = elements.find(qn("a:fontScheme"))
    return clr_scheme, font_scheme


class _FakeMaster:
    """Minimal stand-in exposing only the ``._element`` attribute."""

    def __init__(self, element):
        self._element = element


class _FakeNoWidthPresentation:
    """Simulates a template whose ``slide_width`` is unset (``None``)."""

    def __init__(self):
        self.slide_width = None


# ------------------------------------------------------------------ _hex


def test_hex_valid_without_hash():
    assert mrp._hex("1a2b3c") == "1A2B3C"


def test_hex_valid_with_hash_and_whitespace():
    assert mrp._hex("  #AbCdEf  ") == "ABCDEF"


def test_hex_empty_uses_default():
    assert mrp._hex("") == "000000"


def test_hex_none_uses_default():
    assert mrp._hex(None) == "000000"


def test_hex_wrong_length_uses_custom_default():
    assert mrp._hex("bad", default="FFFFFF") == "FFFFFF"
    assert mrp._hex("1234567", default="FFFFFF") == "FFFFFF"


# ---------------------------------------------------------- _first_font


def test_first_font_empty_uses_default():
    assert mrp._first_font("", "Fallback") == "Fallback"


def test_first_font_single_name():
    assert mrp._first_font("Arial", "Fallback") == "Arial"


def test_first_font_picks_first_of_list_and_strips_quotes():
    value = '"Times New Roman", Arial, sans-serif'
    assert mrp._first_font(value, "Fallback") == "Times New Roman"


def test_first_font_blank_first_entry_uses_default():
    assert mrp._first_font(" , Arial", "Fallback") == "Fallback"


# ------------------------------------------------------ _set_scheme_color


def test_set_scheme_color_replaces_sys_clr_with_srgb():
    clr_scheme, _ = _fresh_theme_elements()
    dk1 = clr_scheme.find(qn("a:dk1"))
    assert dk1.find(qn("a:sysClr")) is not None  # sanity: template default

    mrp._set_scheme_color(clr_scheme, "dk1", "ABCDEF")

    assert dk1.find(qn("a:sysClr")) is None
    srgb = dk1.find(qn("a:srgbClr"))
    assert srgb is not None
    assert srgb.get("val") == "ABCDEF"
    assert len(list(dk1)) == 1  # the old child was removed, not kept


def test_set_scheme_color_replaces_existing_srgb():
    clr_scheme, _ = _fresh_theme_elements()

    mrp._set_scheme_color(clr_scheme, "accent1", "112233")

    accent1 = clr_scheme.find(qn("a:accent1"))
    srgb = accent1.find(qn("a:srgbClr"))
    assert srgb.get("val") == "112233"
    assert len(list(accent1)) == 1


def test_set_scheme_color_unknown_slot_is_noop():
    clr_scheme, _ = _fresh_theme_elements()
    before = etree.tostring(clr_scheme)

    mrp._set_scheme_color(clr_scheme, "not-a-real-slot", "FFFFFF")

    assert etree.tostring(clr_scheme) == before


# -------------------------------------------------------------- _set_font


def test_set_font_updates_latin_typeface():
    _, font_scheme = _fresh_theme_elements()

    mrp._set_font(font_scheme, "majorFont", "Custom Major")

    latin = font_scheme.find(qn("a:majorFont")).find(qn("a:latin"))
    assert latin.get("typeface") == "Custom Major"


def test_set_font_unknown_which_is_noop():
    _, font_scheme = _fresh_theme_elements()
    before = etree.tostring(font_scheme)

    mrp._set_font(font_scheme, "bogusFont", "X")

    assert etree.tostring(font_scheme) == before


def test_set_font_family_without_latin_is_noop():
    root = etree.Element(qn("a:fontScheme"))
    family = etree.SubElement(root, qn("a:majorFont"))
    # No a:latin child: simulates a font-scheme entry missing the
    # typeface element the function expects to update.

    mrp._set_font(root, "majorFont", "Should Not Apply")

    assert family.find(qn("a:latin")) is None
    assert list(family) == []


# ------------------------------------------------------ _widen_placeholders


def test_widen_placeholders_scales_explicit_xfrm_on_master():
    prs = Presentation()
    master = prs.slide_masters[0]
    ratio = float(mrp._WIDESCREEN_W) / float(prs.slide_width)

    before = {}
    for ph in master.placeholders:
        xfrm = ph._element.spPr.find(qn("a:xfrm"))
        off = xfrm.find(qn("a:off"))
        ext = xfrm.find(qn("a:ext"))
        before[ph.placeholder_format.idx] = (
            int(off.get("x")),
            int(ext.get("cx")),
        )

    mrp._widen_placeholders(master, ratio)

    for ph in master.placeholders:
        xfrm = ph._element.spPr.find(qn("a:xfrm"))
        off = xfrm.find(qn("a:off"))
        ext = xfrm.find(qn("a:ext"))
        orig_x, orig_cx = before[ph.placeholder_format.idx]
        assert int(off.get("x")) == int(orig_x * ratio)
        assert int(ext.get("cx")) == int(orig_cx * ratio)


def test_widen_placeholders_layout_mixed_geometry():
    """Layout placeholders mix explicit and inherited geometry.

    ``dt``/``ftr``/``sldNum`` on the default layout inherit their
    ``a:xfrm`` from the master (no explicit element of their own) and
    must be skipped; the title/subtitle placeholders own an explicit
    ``a:xfrm`` and must be scaled.
    """
    prs = Presentation()
    layout = prs.slide_masters[0].slide_layouts[0]
    ratio = 1.5

    before = {}
    for ph in layout.placeholders:
        xfrm = ph._element.spPr.find(qn("a:xfrm"))
        if xfrm is None:
            before[ph.placeholder_format.idx] = None
        else:
            off = xfrm.find(qn("a:off"))
            ext = xfrm.find(qn("a:ext"))
            before[ph.placeholder_format.idx] = (
                int(off.get("x")),
                int(ext.get("cx")),
            )
    assert None in before.values()  # sanity: some inherit geometry
    assert any(v is not None for v in before.values())  # some are explicit

    mrp._widen_placeholders(layout, ratio)

    for ph in layout.placeholders:
        xfrm = ph._element.spPr.find(qn("a:xfrm"))
        idx = ph.placeholder_format.idx
        if before[idx] is None:
            assert xfrm is None  # still inherited, never gained one
        else:
            off = xfrm.find(qn("a:off"))
            ext = xfrm.find(qn("a:ext"))
            orig_x, orig_cx = before[idx]
            assert int(off.get("x")) == int(orig_x * ratio)
            assert int(ext.get("cx")) == int(orig_cx * ratio)


# ------------------------------------------------------ _retune_text_styles


def test_retune_text_styles_updates_real_master():
    prs = Presentation()
    master = prs.slide_masters[0]

    mrp._retune_text_styles(master)

    tx_styles = master._element.find(qn("p:txStyles"))
    title_style = tx_styles.find(qn("p:titleStyle"))
    for pr in title_style.iter(qn("a:defRPr")):
        assert pr.get("sz") == str(mrp._TITLE_SIZE)

    body_style = tx_styles.find(qn("p:bodyStyle"))
    for lvl, size in enumerate(mrp._BODY_SIZES, start=1):
        pr = body_style.find(qn(f"a:lvl{lvl}pPr")).find(qn("a:defRPr"))
        assert pr.get("sz") == str(size)


def test_retune_text_styles_no_tx_styles_is_noop():
    root = etree.Element(qn("p:sldMaster"))  # no p:txStyles at all
    fake = _FakeMaster(root)

    mrp._retune_text_styles(fake)  # must not raise

    assert root.find(qn("p:txStyles")) is None


def test_retune_text_styles_missing_and_malformed_levels():
    """Cover the per-level guards with a deliberately partial master.

    lvl1 has a normal ``sz`` (must be updated); lvl2's ``defRPr`` has
    no ``sz`` attribute (the guard must skip it rather than inventing
    one); lvl3 has no ``defRPr`` at all; lvl4-9 have no ``pPr`` parent
    at all (the ``continue`` branch). ``titleStyle`` is absent too,
    exercising that guard in the same pass.
    """
    root = etree.Element(qn("p:sldMaster"))
    tx_styles = etree.SubElement(root, qn("p:txStyles"))
    body_style = etree.SubElement(tx_styles, qn("p:bodyStyle"))

    lvl1 = etree.SubElement(body_style, qn("a:lvl1pPr"))
    def_rpr1 = etree.SubElement(lvl1, qn("a:defRPr"))
    def_rpr1.set("sz", "1800")

    lvl2 = etree.SubElement(body_style, qn("a:lvl2pPr"))
    etree.SubElement(lvl2, qn("a:defRPr"))  # no sz attribute

    etree.SubElement(body_style, qn("a:lvl3pPr"))  # no defRPr child
    # lvl4pPr..lvl9pPr intentionally missing entirely.

    fake = _FakeMaster(root)

    mrp._retune_text_styles(fake)  # must not raise

    assert def_rpr1.get("sz") == str(mrp._BODY_SIZES[0])
    assert lvl2.find(qn("a:defRPr")).get("sz") is None


# ---------------------------------------------------------- _enable_autofit


def test_enable_autofit_replaces_existing_and_skips_chrome():
    prs = Presentation()
    master = prs.slide_masters[0]

    # Seed the body placeholder (idx 1) with a *different* autofit tag
    # to prove the function replaces it rather than merely appending.
    body_ph = master.placeholders[1]
    body_pr = body_ph.text_frame._txBody.bodyPr
    for old in body_pr.findall(qn("a:normAutofit")):
        body_pr.remove(old)
    etree.SubElement(body_pr, qn("a:noAutofit"))

    mrp._enable_autofit(master)

    for idx in (0, 1):  # title, body: not chrome, must get normAutofit
        bp = master.placeholders[idx].text_frame._txBody.bodyPr
        tags = [etree.QName(c).localname for c in bp]
        assert tags == ["normAutofit"]

    for idx in (2, 3, 4):  # dt, ftr, sldNum: chrome, must be untouched
        bp = master.placeholders[idx].text_frame._txBody.bodyPr
        assert list(bp) == []


# ------------------------------------------------------ _drop_stale_size_type


def test_drop_stale_size_type_removes_type_attr():
    prs = Presentation()
    sld_sz = prs._element.find(qn("p:sldSz"))
    assert sld_sz.get("type") == "screen4x3"  # sanity: template default

    mrp._drop_stale_size_type(prs)

    assert sld_sz.get("type") is None
    assert sld_sz.get("cx") == "9144000"
    assert sld_sz.get("cy") == "6858000"


def test_drop_stale_size_type_noop_when_no_type_attr():
    prs = Presentation()
    sld_sz = prs._element.find(qn("p:sldSz"))
    del sld_sz.attrib["type"]

    mrp._drop_stale_size_type(prs)  # must not raise

    assert sld_sz.get("type") is None


# --------------------------------------------------------- build_reference


def test_build_reference_widescreen_applies_theme(tmp_path):
    theme = _make_theme()
    target = tmp_path / "theme.pptx"

    mrp.build_reference(theme, target, widescreen=True)

    assert target.exists()
    prs = Presentation(str(target))  # really loadable

    assert prs.slide_width == mrp._WIDESCREEN_W
    assert prs.slide_height == mrp._WIDESCREEN_H
    sld_sz = prs._element.find(qn("p:sldSz"))
    assert sld_sz.get("type") is None  # stale "screen4x3" dropped

    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    theme_el = etree.fromstring(theme_part.blob)
    elements = theme_el.find(qn("a:themeElements"))
    clr_scheme = elements.find(qn("a:clrScheme"))
    font_scheme = elements.find(qn("a:fontScheme"))

    css = theme.css_vars
    for slot, var in mrp._COLOR_MAP.items():
        srgb = clr_scheme.find(qn(f"a:{slot}")).find(qn("a:srgbClr"))
        assert srgb.get("val") == mrp._hex(css[var])

    major = font_scheme.find(qn("a:majorFont")).find(qn("a:latin"))
    minor = font_scheme.find(qn("a:minorFont")).find(qn("a:latin"))
    assert major.get("typeface") == "Header Font"
    assert minor.get("typeface") == "Body Font"

    tx_styles = master._element.find(qn("p:txStyles"))
    title_style = tx_styles.find(qn("p:titleStyle"))
    for pr in title_style.iter(qn("a:defRPr")):
        assert pr.get("sz") == str(mrp._TITLE_SIZE)
    body_style = tx_styles.find(qn("p:bodyStyle"))
    for lvl, size in enumerate(mrp._BODY_SIZES, start=1):
        pr = body_style.find(qn(f"a:lvl{lvl}pPr")).find(qn("a:defRPr"))
        assert pr.get("sz") == str(size)

    title_ph = master.placeholders[0]
    body_pr = title_ph.text_frame._txBody.bodyPr
    assert body_pr.find(qn("a:normAutofit")) is not None

    default_prs = Presentation()
    ratio = float(mrp._WIDESCREEN_W) / float(default_prs.slide_width)
    default_title = default_prs.slide_masters[0].placeholders[0]
    orig_off = default_title._element.spPr.find(qn("a:xfrm")).find(
        qn("a:off")
    )
    new_off = title_ph._element.spPr.find(qn("a:xfrm")).find(qn("a:off"))
    assert int(new_off.get("x")) == int(int(orig_off.get("x")) * ratio)


def test_build_reference_no_widescreen_keeps_native_canvas(tmp_path):
    theme = _make_theme()
    target = tmp_path / "theme_43.pptx"

    mrp.build_reference(theme, target, widescreen=False)

    prs = Presentation(str(target))
    default_prs = Presentation()
    assert prs.slide_width == default_prs.slide_width
    assert prs.slide_height == default_prs.slide_height

    sld_sz = prs._element.find(qn("p:sldSz"))
    assert sld_sz.get("type") == "screen4x3"  # never touched

    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    theme_el = etree.fromstring(theme_part.blob)
    clr_scheme = theme_el.find(qn("a:themeElements")).find(
        qn("a:clrScheme")
    )
    dk1_srgb = clr_scheme.find(qn("a:dk1")).find(qn("a:srgbClr"))
    assert dk1_srgb.get("val") == mrp._hex(theme.css_vars["fg"])


def test_build_reference_raises_when_template_has_no_slide_width(
    monkeypatch, tmp_path,
):
    monkeypatch.setattr(mrp, "Presentation", _FakeNoWidthPresentation)
    theme = _make_theme()

    with pytest.raises(ValueError, match="cannot be widened to 16:9"):
        mrp.build_reference(theme, tmp_path / "out.pptx", widescreen=True)


# --------------------------------------------------------------------- main


def test_main_writes_all_themes_and_creates_init(tmp_path, monkeypatch):
    fake_out = tmp_path / "out"
    monkeypatch.setattr(mrp, "ROOT", tmp_path)
    monkeypatch.setattr(mrp, "OUT_DIR", fake_out)

    result = mrp.main()

    assert result == 0
    init_file = fake_out / "__init__.py"
    assert init_file.exists()
    assert "Per-theme PowerPoint reference decks" in init_file.read_text(
        encoding="utf-8",
    )

    assert mrp.themes.THEMES  # sanity: bundled themes are loaded
    for theme_id in mrp.themes.THEMES:
        wide = fake_out / f"{theme_id}.pptx"
        narrow = fake_out / f"{theme_id}_43.pptx"
        assert wide.exists()
        assert narrow.exists()
        Presentation(str(wide))  # really loadable
        Presentation(str(narrow))


def test_main_does_not_overwrite_existing_init(tmp_path, monkeypatch):
    fake_out = tmp_path / "out"
    fake_out.mkdir(parents=True)
    marker = "# custom marker, must survive\n"
    (fake_out / "__init__.py").write_text(marker, encoding="utf-8")

    monkeypatch.setattr(mrp, "ROOT", tmp_path)
    monkeypatch.setattr(mrp, "OUT_DIR", fake_out)

    result = mrp.main()

    assert result == 0
    assert (fake_out / "__init__.py").read_text(encoding="utf-8") == marker


def test_dunder_main_guard_exits_zero_without_writing_to_repo(monkeypatch):
    """Exercise ``if __name__ == "__main__": raise SystemExit(main())``.

    ``runpy.run_path`` re-executes the file fresh, so ``ROOT``/
    ``OUT_DIR`` are recomputed from the *real* ``__file__`` and would
    point at the real bundled asset tree — patching them from here
    would not reach that fresh execution anyway, since it rebuilds
    its own globals from scratch. Instead, ``themes.THEMES`` (a
    dict on the already-imported, shared ``themes`` module) is
    emptied so ``main``'s loop runs zero iterations: ``OUT_DIR.mkdir``
    and the ``__init__.py`` check become no-ops against the
    already-existing real directory, and no ``build_reference`` call
    (hence no ``.pptx`` write) ever happens. The before/after snapshot
    below is the actual proof that nothing under ``src/`` moved.
    """
    module_file = Path(mrp.__file__)
    real_out_dir = mrp.OUT_DIR
    before = {p.name: p.stat().st_mtime_ns for p in real_out_dir.iterdir()}

    monkeypatch.setattr(mrp.themes, "THEMES", {})

    with pytest.raises(SystemExit) as exc_info:
        runpy.run_path(str(module_file), run_name="__main__")

    assert exc_info.value.code == 0

    after = {p.name: p.stat().st_mtime_ns for p in real_out_dir.iterdir()}
    assert after == before
